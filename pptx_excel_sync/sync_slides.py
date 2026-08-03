"""Sync a PowerPoint deck against an updated Excel sheet.

Given a "previous" ``.pptx`` deck and an updated ``.xlsx`` sheet, this
module produces a new deck where:

- Slides that already exist for a spreadsheet row (matched via a hidden
  ``KEY:`` marker in the slide's speaker notes) are refreshed in place.
- Spreadsheet rows with no matching slide yet get a new slide, cloned
  from a "template" slide (marked with ``TEMPLATE: <name>`` in its
  speaker notes) so the new slide matches the deck's existing style.

See ``README.md`` in this directory for the full spreadsheet/placeholder
format and a runnable example.
"""

from __future__ import annotations

import argparse
import copy
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
TEMPLATE_MARKER_RE = re.compile(r"^TEMPLATE:\s*(.*)$", re.MULTILINE)
KEY_MARKER_RE = re.compile(r"^KEY:\s*(.*)$", re.MULTILINE)
TEMPLATE_SOURCE_MARKER_RE = re.compile(r"^TEMPLATE_SOURCE:\s*(.*)$", re.MULTILINE)


def _normalize(name: object) -> str:
    """Normalize a header/placeholder name so ``Status``, ``status`` and
    ``{{ status }}`` all refer to the same field."""
    return re.sub(r"\s+", "_", str(name).strip().lower())


@dataclass
class SyncResult:
    updated: list[str] = field(default_factory=list)
    created: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> str:
        lines = [
            f"Updated {len(self.updated)} slide(s): {', '.join(self.updated) or '-'}",
            f"Created {len(self.created)} slide(s): {', '.join(self.created) or '-'}",
        ]
        if self.skipped:
            lines.append(f"Skipped {len(self.skipped)} row(s): {', '.join(self.skipped)}")
        for warning in self.warnings:
            lines.append(f"WARNING: {warning}")
        return "\n".join(lines)


def load_rows(excel_path: str, sheet_name: Optional[str] = None) -> list[dict]:
    """Read spreadsheet rows into normalized dicts, keyed by normalized
    header name. Requires a ``key`` column (case/whitespace-insensitive).
    """
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]

    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []

    headers = [_normalize(cell) if cell is not None else "" for cell in header_row]
    if "key" not in headers:
        raise ValueError(
            "Excel sheet must have a 'key' column in its header row "
            f"(found columns: {[h for h in headers if h]})"
        )

    rows = []
    for raw in rows_iter:
        if raw is None or all(value is None for value in raw):
            continue
        row = {
            header: value
            for header, value in zip(headers, raw)
            if header
        }
        key = row.get("key")
        if key is None or str(key).strip() == "":
            continue
        row["key"] = str(key).strip()
        rows.append(row)
    return rows


def _get_notes_text(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text or ""


def _find_slide_by_key(prs, key: str):
    for slide in prs.slides:
        match = KEY_MARKER_RE.search(_get_notes_text(slide))
        if match and match.group(1).strip() == key:
            return slide
    return None


def _template_name_of(slide) -> Optional[str]:
    match = TEMPLATE_MARKER_RE.search(_get_notes_text(slide))
    return match.group(1).strip() if match else None


def _find_template_slide(prs, name: Optional[str] = None):
    fallback = None
    for slide in prs.slides:
        slide_name = _template_name_of(slide)
        if slide_name is None:
            continue
        if fallback is None:
            fallback = slide
        if name and slide_name == name:
            return slide
    return None if name else fallback


def _copy_shapes(source_slide, target_slide) -> None:
    """Replace ``target_slide``'s shapes with a shape-for-shape copy of
    ``source_slide``'s shapes.

    Note: shapes that reference external media via relationships (e.g.
    pictures, embedded chart workbooks) are copied by XML only, so this
    works best for text- and table-based template slides.
    """
    for shape in list(target_slide.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in source_slide.shapes:
        target_slide.shapes._spTree.append(copy.deepcopy(shape._element))


def _new_slide_from_template(prs, template_slide):
    """Append a brand-new slide cloned from ``template_slide``, using the
    same layout, and return it. Notes are set separately by the caller."""
    new_slide = prs.slides.add_slide(template_slide.slide_layout)
    _copy_shapes(template_slide, new_slide)
    return new_slide


def _set_generated_notes(slide, key: str, template_name: Optional[str]) -> None:
    lines = [f"KEY: {key}"]
    if template_name:
        lines.append(f"TEMPLATE_SOURCE: {template_name}")
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def _substitute(text: str, data: dict, missing: set[str]) -> str:
    def repl(match: re.Match) -> str:
        field_name = match.group(1)
        normalized = _normalize(field_name)
        if normalized in data:
            value = data[normalized]
            return "" if value is None else str(value)
        missing.add(field_name.strip())
        return match.group(0)

    return PLACEHOLDER_RE.sub(repl, text)


def _fill_text_frame(text_frame, data: dict, missing: set[str]) -> None:
    for paragraph in text_frame.paragraphs:
        original = "".join(run.text for run in paragraph.runs)
        if "{{" not in original:
            continue
        updated = _substitute(original, data, missing)
        if updated == original:
            continue
        if paragraph.runs:
            paragraph.runs[0].text = updated
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.text = updated


def _fill_shape(shape, data: dict, missing: set[str]) -> None:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub_shape in shape.shapes:
            _fill_shape(sub_shape, data, missing)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                _fill_text_frame(cell.text_frame, data, missing)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        _fill_text_frame(shape.text_frame, data, missing)


def _fill_slide(slide, data: dict, missing: set[str]) -> None:
    for shape in slide.shapes:
        _fill_shape(shape, data, missing)


def sync_deck(
    previous_path: str,
    excel_path: str,
    output_path: Optional[str] = None,
    sheet_name: Optional[str] = None,
    dry_run: bool = False,
) -> SyncResult:
    """Build a new deck from ``previous_path`` fixed/extended with rows
    from ``excel_path``, and write it to ``output_path`` (unless
    ``dry_run`` is set, in which case nothing is written)."""
    prs = Presentation(previous_path)
    rows = load_rows(excel_path, sheet_name)
    result = SyncResult()

    for row in rows:
        key = row["key"]
        data = {k: v for k, v in row.items() if k not in ("key", "template")}
        missing: set[str] = set()

        requested_template = row.get("template")
        requested_template = str(requested_template).strip() if requested_template else None

        slide = _find_slide_by_key(prs, key)
        if slide is not None:
            # Re-derive the slide's shapes from its original blueprint before
            # substituting, so placeholders consumed by an earlier run are
            # available again. Without a recorded blueprint (e.g. a
            # hand-authored slide), fall back to substituting in place.
            source_name = TEMPLATE_SOURCE_MARKER_RE.search(_get_notes_text(slide))
            source_name = source_name.group(1).strip() if source_name else None
            if source_name:
                source_template = _find_template_slide(prs, source_name)
                if source_template is not None:
                    _copy_shapes(source_template, slide)
                else:
                    result.warnings.append(
                        f"Row '{key}': original template '{source_name}' no longer "
                        "found in the deck; updating existing slide content as-is"
                    )
            _fill_slide(slide, data, missing)
            _set_generated_notes(slide, key, source_name or requested_template)
            result.updated.append(key)
        else:
            template_slide = _find_template_slide(prs, requested_template)
            if template_slide is None:
                reason = (
                    f"no slide with TEMPLATE: {requested_template}"
                    if requested_template
                    else "no TEMPLATE slide found in the deck"
                )
                result.warnings.append(f"Row '{key}': {reason}; skipped")
                result.skipped.append(key)
                continue
            resolved_template_name = _template_name_of(template_slide) or requested_template
            new_slide = _new_slide_from_template(prs, template_slide)
            _fill_slide(new_slide, data, missing)
            _set_generated_notes(new_slide, key, resolved_template_name)
            result.created.append(key)

        for field_name in sorted(missing):
            result.warnings.append(
                f"Row '{key}': no spreadsheet column for placeholder '{{{{{field_name}}}}}'"
            )

    if not dry_run:
        if output_path is None:
            raise ValueError("output_path is required unless dry_run=True")
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)

    return result


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--previous", required=True, help="Path to the previous .pptx deck")
    parser.add_argument("--excel", required=True, help="Path to the updated .xlsx sheet")
    parser.add_argument("--output", help="Path to write the new .pptx deck")
    parser.add_argument("--sheet", help="Worksheet name to read (defaults to the first sheet)")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Report what would change without writing an output file",
    )
    args = parser.parse_args(argv)

    if not args.dry_run and not args.output:
        parser.error("--output is required unless --dry-run is set")

    result = sync_deck(
        previous_path=args.previous,
        excel_path=args.excel,
        output_path=args.output,
        sheet_name=args.sheet,
        dry_run=args.dry_run,
    )
    print(result.summary())
    return 0


if __name__ == "__main__":
    sys.exit(main())
