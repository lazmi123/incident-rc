"""Generate a small demo deck + spreadsheet for trying out sync_slides.py.

Run:

    python pptx_excel_sync/examples/generate_samples.py

This writes ``previous_deck.pptx`` and ``updates.xlsx`` next to this
script.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from pptx import Presentation

HERE = Path(__file__).parent


def build_previous_deck(path: Path) -> None:
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    title_slide = prs.slides.add_slide(title_layout)
    title_slide.shapes.title.text = "Weekly Incident Review"
    title_slide.placeholders[1].text = "Generated from the incident tracker"

    # Reusable blueprint for incident-status slides. Marked TEMPLATE so
    # sync_slides.py knows to clone it for new spreadsheet rows.
    template_slide = prs.slides.add_slide(content_layout)
    template_slide.shapes.title.text = "{{title}}"
    body = template_slide.placeholders[1].text_frame
    body.text = "Status: {{status}}"
    body.add_paragraph().text = "Owner: {{owner}}"
    body.add_paragraph().text = "Impact: {{impact}}"
    template_slide.notes_slide.notes_text_frame.text = "TEMPLATE: incident-status"

    # A slide already generated on a previous run for INC-001, with a
    # deliberately stale status so the example shows the "fix" behavior.
    inc001_slide = prs.slides.add_slide(content_layout)
    inc001_slide.shapes.title.text = "INC-001: Checkout latency spike"
    body = inc001_slide.placeholders[1].text_frame
    body.text = "Status: Investigating"
    body.add_paragraph().text = "Owner: Priya"
    body.add_paragraph().text = "Impact: 5% of checkouts delayed"
    inc001_slide.notes_slide.notes_text_frame.text = (
        "KEY: INC-001\nTEMPLATE_SOURCE: incident-status"
    )

    prs.save(path)


def build_updates_sheet(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Incidents"
    ws.append(["key", "template", "title", "status", "owner", "impact"])
    # INC-001 already has a slide (KEY: INC-001) -> gets updated in place.
    ws.append([
        "INC-001",
        "incident-status",
        "INC-001: Checkout latency spike",
        "Resolved",
        "Priya",
        "5% of checkouts delayed for 40 minutes",
    ])
    # INC-002 has no existing slide -> a new one is cloned from the template.
    ws.append([
        "INC-002",
        "incident-status",
        "INC-002: Login API 500s",
        "Mitigated",
        "Jordan",
        "1.2% of logins failed for 12 minutes",
    ])
    wb.save(path)


def main() -> None:
    build_previous_deck(HERE / "previous_deck.pptx")
    build_updates_sheet(HERE / "updates.xlsx")
    print(f"Wrote {HERE / 'previous_deck.pptx'}")
    print(f"Wrote {HERE / 'updates.xlsx'}")


if __name__ == "__main__":
    main()
